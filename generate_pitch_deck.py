import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(11, 17, 30)      # #0b111e
    CARD_BG = RGBColor(19, 28, 46)      # #131c2e
    ACCENT_GOLD = RGBColor(245, 158, 11)# #f59e0b
    ACCENT_BLUE = RGBColor(59, 130, 246)# #3b82f6
    TEXT_WHITE = RGBColor(241, 245, 249)# #f1f5f9
    TEXT_MUTED = RGBColor(148, 163, 184)# #94a3b8
    GREEN_COLOR = RGBColor(16, 185, 129)# #10b981
    RED_COLOR = RGBColor(239, 68, 68)   # #ef4444

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, subtitle_text="SANGYAN 2026: SEBI & NSDL Investor Resilience Hackathon"):
        # Header Badge / Ribbon
        badge = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        p_b.text = f"🛡️ {subtitle_text.upper()}"
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = ACCENT_GOLD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Ribbon Tag
    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(5.8), Inches(0.45))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = CARD_BG
    badge1.line.color.rgb = ACCENT_GOLD
    tf1 = badge1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "SANGYAN 2026 • SEBI | NSDL | IIT (BHU) VARANASI"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_GOLD

    # Main Title
    t1_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.2))
    tf1_t = t1_box.text_frame
    tf1_t.word_wrap = True
    p1_t = tf1_t.paragraphs[0]
    p1_t.text = "ScamSatark (स्कैम सतर्क)"
    p1_t.font.size = Pt(46)
    p1_t.font.bold = True
    p1_t.font.color.rgb = TEXT_WHITE

    p1_sub = tf1_t.add_paragraph()
    p1_sub.text = "AI-Powered National Investor Resilience & Digital Fraud Interception Infrastructure"
    p1_sub.font.size = Pt(22)
    p1_sub.font.color.rgb = ACCENT_GOLD

    # Subtext / Core premise
    desc1_box = s1.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.7), Inches(1.5))
    tf1_d = desc1_box.text_frame
    tf1_d.word_wrap = True
    p1_d = tf1_d.paragraphs[0]
    p1_d.text = "A zero-login, public-good defense platform protecting first-time & Tier-2/3 retail investors from deceptive financial vectors, fake SEBI claims, and pump-and-dump channels in 22 Official Indian Languages."
    p1_d.font.size = Pt(15)
    p1_d.font.color.rgb = TEXT_MUTED

    # Highlights Row
    for i, (metric, label) in enumerate([
        ("22 Languages", "Full 8th Schedule Native Support"),
        ("< 5ms Engine", "Deterministic Heuristics + Gemini Multimodal"),
        ("Public Good", "Strict Non-Commercial Guardrail Compliance"),
        ("1-Click Redressal", "Direct SCORES & Cybercrime 1930 Dossier")
    ]):
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*2.95), Inches(5.8), Inches(2.8), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE if i%2==0 else ACCENT_GOLD
        tf = card.text_frame
        tf.word_wrap = True
        p_m = tf.paragraphs[0]
        p_m.text = metric
        p_m.font.size = Pt(17)
        p_m.font.bold = True
        p_m.font.color.rgb = TEXT_WHITE
        p_l = tf.add_paragraph()
        p_l.text = label
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 2: The Challenge & Ground Reality
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Challenge: Rapid Market Access Outrunning Financial Confidence")

    cards_data_2 = [
        ("📈 16+ Crore Demat Surge", "Over 70% incremental retail account growth originates from non-metro, Tier-2, and Tier-3 cities. Emerging investors lack institutional guardrails.", RED_COLOR),
        ("⚠️ 9 Out of 10 F&O Loss Rate", "SEBI's landmark study revealed 90% of individual derivative traders lose money. Market democratization without resilience leads to catastrophic capital erosion.", ACCENT_GOLD),
        ("🎭 Digital Fraud Vector Epidemic", "Scammers weaponize fake SEBI IDs, Telegram pump-and-dump tip groups, cloned broker portals, remote APKs (AnyDesk), and synthetic urgency.", ACCENT_BLUE)
    ]

    for i, (title, text, color) in enumerate(cards_data_2):
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.95), Inches(1.8), Inches(3.8), Inches(3.6))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = color
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = f"\n{text}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_MUTED

    banner_2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.7), Inches(11.7), Inches(1.1))
    banner_2.fill.solid()
    banner_2.fill.fore_color.rgb = CARD_BG
    banner_2.line.color.rgb = ACCENT_GOLD
    tf_b2 = banner_2.text_frame
    tf_b2.word_wrap = True
    p_b2 = tf_b2.paragraphs[0]
    p_b2.text = "🎯 Core SANGYAN Mandate:"
    p_b2.font.bold = True
    p_b2.font.size = Pt(14)
    p_b2.font.color.rgb = ACCENT_GOLD
    p_b2_sub = tf_b2.add_paragraph()
    p_b2_sub.text = "The objective is NOT to give investment tips. It is to build resilient investors who can spot manipulation, avoid loss, understand financial risks, and decide safely."
    p_b2_sub.font.size = Pt(13)
    p_b2_sub.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 3: ScamSatark Solution Architecture
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "ScamSatark: A 360° Investor Protection Shield")

    features_s3 = [
        ("🔍 Real-Time Vector Verifier", "Paste WhatsApp/Telegram tip forwards or links. Immediate sub-5ms heuristic triage across 15+ known fraud patterns.", ACCENT_GOLD),
        ("📷 Multimodal Screenshot OCR", "Upload screenshot of chats, fake SEBI certificates, or cloned portals. Gemini AI extracts text, logos, and hidden traps.", ACCENT_BLUE),
        ("🗣️ Dual-Engine Voice In & Out", "Speak in your native tongue via Web Speech / Gemini audio. Hear instant audio verdicts in authentic regional pronunciation.", GREEN_COLOR),
        ("🧠 Cognitive Manipulation Deconstruction", "Explains the psychological bait (Fake Urgency, Greed, Impersonation) using 8th-grade everyday analogies.", ACCENT_GOLD),
        ("📋 1-Click SCORES Redressal Dossier", "Automatically generates a structured incident report ready to file on SEBI SCORES 2.0 or National Cybercrime Helpline 1930.", ACCENT_BLUE),
        ("🌐 22 Official Indian Languages", "Pre-bundled offline dictionary covering all 22 8th Schedule languages with instant zero-latency language switching.", GREEN_COLOR)
    ]

    for i, (title, desc, color) in enumerate(features_s3):
        col = i % 3
        row = i // 3
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col*3.95), Inches(1.8 + row*2.6), Inches(3.8), Inches(2.3))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = color
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 4: Track Alignment & Impact
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Direct SANGYAN Track Alignment & Evaluation Rubric")

    tracks = [
        ("Track A: Digital Fraud & Scam Resilience", "PRIMARY FOCUS", "• Intercepts deceptive vectors before money changes hands.\n• Scans fake SEBI numbers, pump-and-dump groups, and remote APKs (AnyDesk).\n• Tip-group risk profiler with calibrated risk scores.", ACCENT_GOLD),
        ("Track B: Grievance & SCORES Integration", "ACTIONABLE REDRESSAL", "• Formats 1-Click structured incident dossiers.\n• Seamless escalation pathways to SEBI SCORES 2.0 & Cybercrime 1930.\n• Guides first-time investors through evidence collection.", ACCENT_BLUE),
        ("Track C & E: Bharat Literacy & Content Literacy", "EDUCATION IN BHARAT", "• Voice-first explainer in 22 regional languages.\n• Replaces jargon with plain analogies.\n• Deconstructs social proof & greed baiting tactics.", GREEN_COLOR)
    ]

    for i, (t_title, badge, points, col_accent) in enumerate(tracks):
        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.95), Inches(1.8), Inches(3.8), Inches(5.0))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col_accent
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = t_title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p_badge = tf.add_paragraph()
        p_badge.text = f"[{badge}]"
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = col_accent
        p2 = tf.add_paragraph()
        p2.text = f"\n{points}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 5: Technical Architecture
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "High-Resilience Hybrid Technical Architecture")

    tech_boxes = [
        ("Layer 1: Deterministic Engine", "Client-Side Regex & Rule Heuristics\n• 15+ calibrated SEBI fraud patterns\n• Sub-5ms instant triage\n• 100% offline functionality", ACCENT_GOLD),
        ("Layer 2: Multimodal Gemini AI", "Gemini 3.5 Flash-Lite + Fallbacks\n• High-throughput multimodal OCR\n• Cognitive manipulation classification\n• Multi-model fallback cascade (Zero 429 errors)", ACCENT_BLUE),
        ("Layer 3: Bharat Voice & i18n", "22 Languages + Web Speech API\n• Continuous client-side recognition\n• Regional SpeechSynthesis with BCP-47 tags\n• Gemini audio transcription fallback", GREEN_COLOR),
        ("Layer 4: Redressal Infrastructure", "Government & Regulatory Integration\n• Direct SEBI SCORES 2.0 template generator\n• National Cybercrime 1930 hotlines\n• Direct SEBI Registry Intermediary verification", ACCENT_GOLD)
    ]

    for i, (title, desc, color) in enumerate(tech_boxes):
        col = i % 2
        row = i // 2
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col*5.95), Inches(1.8 + row*2.6), Inches(5.75), Inches(2.3))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = color
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 6: Bharat-First 22 Languages
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Tier-2/3 Usability: Bharat-First Inclusion in 22 Languages")

    box_lang = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(3.2))
    box_lang.fill.solid()
    box_lang.fill.fore_color.rgb = CARD_BG
    box_lang.line.color.rgb = GREEN_COLOR
    tf_l = box_lang.text_frame
    tf_l.word_wrap = True
    p_l1 = tf_l.paragraphs[0]
    p_l1.text = "🌍 Comprehensive 8th Schedule Language Matrix (Zero-Latency Pre-Bundled Dictionaries):"
    p_l1.font.size = Pt(16)
    p_l1.font.bold = True
    p_l1.font.color.rgb = GREEN_COLOR

    p_l2 = tf_l.add_paragraph()
    p_l2.text = "\n• Hindi (हिन्दी)  • Marathi (मराठी)  • Bengali (বাংলা)  • Telugu (తెలుగు)  • Tamil (தமிழ்)\n• Gujarati (ગુજરાતી)  • Kannada (ಕನ್ನಡ)  • Malayalam (മലയാളം)  • Punjabi (ਪੰਜਾਬੀ)  • Odia (ଓଡ଼ିଆ)\n• Assamese (অসমীয়া)  • Urdu (اردو)  • Sanskrit (संस्कृतम्)  • Maithili (मैथिली)  • Santali (ᱥᱟᱱᱛᱟᱲᱤ)\n• Kashmiri (کٲشُر)  • Nepali (नेपाली)  • Konkani (कोंकणी)  • Sindhi (سنڌي)  • Dogri (डोगरी)\n• Manipuri (মৈতৈলোন্)  • Bodo (बड़ो)  • English"
    p_l2.font.size = Pt(14)
    p_l2.font.color.rgb = TEXT_WHITE

    # Key Bharat features
    bharat_points = [
        ("0ms Instant Switch", "Language changes dynamically without network roundtrips.", ACCENT_GOLD),
        ("Voice-In & Voice-Out", "Illiterate and senior citizens can speak and listen to warnings.", ACCENT_BLUE),
        ("Ultra Low-Bandwidth", "Lightweight single-page payload (<50KB compressed). Runs on 2G/3G.", GREEN_COLOR)
    ]
    for i, (t, d, col) in enumerate(bharat_points):
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.95), Inches(5.3), Inches(3.8), Inches(1.6))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = t
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 7: Mandatory Guardrail Compliance
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Strict Guardrail Compliance: A Pure Public-Good Infrastructure")

    guardrails = [
        ("🚫 Zero Commercial / Speculative Outcome", "Strictly NO stock tips, buy/sell/hold calls, price predictions, or broker promotions. Programmatic guardrail regex blocks any generated recommendation.", RED_COLOR),
        ("🚫 Zero Monetization Funnels", "No broking commissions, margin nudges, paid subscriptions, or sponsor upsells. Operates 100% as public-benefit civic infrastructure.", ACCENT_GOLD),
        ("🛡️ Privacy by Design", "No mandatory login. No harvesting of private SMS, OTPs, contacts, or financial accounts. Client-side evaluation ensures confidential checks.", GREEN_COLOR),
        ("⚖️ Transparent Uncertainty", "Never provides false reassurance or claims an unverified tip is '100% safe'. Clearly communicates calibrated risk warnings.", ACCENT_BLUE)
    ]

    for i, (title, desc, color) in enumerate(guardrails):
        col = i % 2
        row = i // 2
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col*5.95), Inches(1.8 + row*2.6), Inches(5.75), Inches(2.3))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = color
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 8: Scalability & Real-World Roadmap
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Scalability, Public Digital Infrastructure & Roadmap")

    roadmap_steps = [
        ("Phase 1: Present Build", "• Real-time scam verifier\n• 22 regional languages\n• Multimodal OCR & Voice\n• SCORES Dossier generator", ACCENT_GOLD),
        ("Phase 2: SEBI & DPI Integration", "• Direct API link with SEBI Intermediary Registry\n• Bhashini AI voice translation pipeline\n• DigiLocker & Account Aggregator tie-in", ACCENT_BLUE),
        ("Phase 3: National Public Utility", "• WhatsApp/Telegram automated helpline bot\n• Telco SMS firewall integration\n• Village Gram Panchayat & SHG awareness circles", GREEN_COLOR)
    ]

    for i, (phase, details, color) in enumerate(roadmap_steps):
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.95), Inches(1.8), Inches(3.8), Inches(5.0))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = color
        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = phase
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = f"\n{details}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 9: Summary & Live Demo Link
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)

    add_header(s9, "Summary: Protecting India's Next 100 Million Investors")

    summary_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    summary_card.fill.solid()
    summary_card.fill.fore_color.rgb = CARD_BG
    summary_card.line.color.rgb = ACCENT_GOLD
    tf_s = summary_card.text_frame
    tf_s.word_wrap = True

    p_s1 = tf_s.paragraphs[0]
    p_s1.text = "🚀 Why ScamSatark Delivers on SANGYAN's Vision:"
    p_s1.font.size = Pt(20)
    p_s1.font.bold = True
    p_s1.font.color.rgb = ACCENT_GOLD

    points_sum = (
        "\n1. High Measurable Impact (30%): Directly stops money from changing hands in predatory scams."
        "\n2. Bharat-First Inclusion (25%): 22 Indian languages, voice-first accessibility, low-bandwidth ready."
        "\n3. 100% Guardrail Compliant (15%): No commercial bias, zero stock recommendations, privacy by design."
        "\n4. Superior Tech Execution (15%): Multi-model fallback cascade, sub-5ms heuristic engine, multimodal vision."
        "\n5. Feasibility & Scalability (15%): Ready for immediate deployment across SEBI, NSDL & Cybercrime portals."
    )
    p_s2 = tf_s.add_paragraph()
    p_s2.text = points_sum
    p_s2.font.size = Pt(15)
    p_s2.font.color.rgb = TEXT_WHITE

    prs.save("c:/ScamSatark/ScamSatark_SANGYAN_Presentation.pptx")
    print("Successfully generated ScamSatark_SANGYAN_Presentation.pptx!")

if __name__ == '__main__':
    create_presentation()
